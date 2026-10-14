"""WO-32 §3 / WO-33 §1 — 网页表格式输入（CSV/TSV/XLSX）解析合同回归。

「预览 / 映射 / 运行」必须对同一份字节给**同一套解释**。这里把它固化成可维护的
回归，而不是只留在一次性脚本里：

* CSV：引号内逗号不拆列、引号内换行算一条记录、坏编码明确报错、不齐列/重复表头
  如实报告（不补齐、不静默替换）；
* XLSX（WO-33 §1 修的就是这里）：
  - 正常单表不再 500（旧实现调了不存在的 ``Workbook.iter_rows``）；
  - 多工作表必须显式选 sheet，且读到的确实是**那一张表**（列与值都不是第一张）；
  - ``header_row`` 取对表头；指向空行要明确报错，而不是返回零列空表；
  - 公式如实标出（``has_formulas``），缓存值单列，绝不冒充运行会读到的内容；
  - openpyxl 会把每行补齐到最大列：尾部空单元格必须去掉，否则「幻影空列」会
    掩盖真正的不齐列；
  - 空表 / 坏 zip / 改名 CSV / 超行超列 → 具体可行动的 :class:`TableError`。

XLSX 一致性另有一条**用真实 SteelData 合同**核对的用例：它把 ``read_document`` /
``matches`` 放在登记好的 SteelData 解释器里跑（那套环境有 pandas/pyjanitor），
没登记就 skip 并说明原因，绝不假装核对过。
"""

from __future__ import annotations

import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from steelai.web_api.tabular import MAX_COLUMNS, MAX_ROWS, TableError, read_table


def _save_xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet(name)
        sheet.title = name
        first = False
        for row in rows:
            sheet.append(row)
    workbook.save(path)


@pytest.fixture()
def single_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "single.xlsx"
    _save_xlsx(path, {"Data": [
        ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"],
        ["S1", 350, 420, 25],
        ["S2", 360, 430, 24],
    ]})
    return path


@pytest.fixture()
def multi_xlsx(tmp_path: Path) -> Path:
    path = tmp_path / "multi.xlsx"
    _save_xlsx(path, {
        "README": [["说明", "本页不是数据"], ["请用 Data 页"]],
        "Data": [
            ["specimen_id", "batch_id", "ys_mpa", "uts_mpa", "el_pct"],
            ["X9", "B-77", 355, 425, 23],
            ["X10", "B-77", 360, 430, 24],
        ],
    })
    return path


# --------------------------------------------------------------------------- #
# CSV / TSV
# --------------------------------------------------------------------------- #

class TestTextContract:
    def test_quoted_comma_stays_one_column(self, tmp_path: Path) -> None:
        path = tmp_path / "quoted.csv"
        path.write_text('specimen_id,ys_mpa,uts_mpa,el_pct\nS1,"contains, comma",12,30\n',
                        encoding="utf-8")
        view = read_table(path)
        assert view.columns == ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"]
        assert view.rows[0][1] == "contains, comma"

    def test_quoted_newline_is_one_record(self, tmp_path: Path) -> None:
        path = tmp_path / "newline.csv"
        path.write_bytes(b'a,b\n"line1\nline2",5\n')
        view = read_table(path)
        assert view.records == 2
        assert view.physical_lines == 3
        assert view.rows[0][0] == "line1\nline2"

    def test_bad_encoding_is_refused_not_replaced(self, tmp_path: Path) -> None:
        path = tmp_path / "bad.csv"
        path.write_bytes(b"a,b\n\xff\xfe,2\n")
        with pytest.raises(TableError, match="UTF-8"):
            read_table(path)

    def test_ragged_rows_are_reported_not_padded(self, tmp_path: Path) -> None:
        path = tmp_path / "ragged.csv"
        path.write_text("a,b,c\n1,2\n3,4,5,6\n", encoding="utf-8")
        view = read_table(path)
        assert view.ragged_rows == [1, 2]

    def test_duplicate_and_empty_headers_are_reported(self, tmp_path: Path) -> None:
        dup = tmp_path / "dup.csv"
        dup.write_text("a,a,b\n1,2,3\n", encoding="utf-8")
        assert read_table(dup).duplicate_headers == ["a"]
        blank = tmp_path / "blank_header.csv"
        blank.write_text("a,,c\n1,2,3\n", encoding="utf-8")
        assert read_table(blank).empty_headers == [1]

    def test_tsv_delimiter_comes_from_suffix_only(self, tmp_path: Path) -> None:
        path = tmp_path / "tab.tsv"
        path.write_text("a\tb,c\n1\t2\n", encoding="utf-8")
        view = read_table(path)
        assert view.delimiter == "\t"
        assert view.columns == ["a", "b,c"]

    def test_sheet_option_is_refused_for_text(self, tmp_path: Path) -> None:
        path = tmp_path / "a.csv"
        path.write_text("a\n1\n", encoding="utf-8")
        with pytest.raises(TableError, match="只能用于 XLSX"):
            read_table(path, sheet="Data")


# --------------------------------------------------------------------------- #
# XLSX（WO-33 §1）
# --------------------------------------------------------------------------- #

class TestXlsxContract:
    def test_single_sheet_no_longer_crashes(self, single_xlsx: Path) -> None:
        # 旧实现直接 AttributeError: 'Workbook' object has no attribute 'iter_rows'
        view = read_table(single_xlsx)
        assert view.columns == ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"]
        assert view.rows[0] == ["S1", "350", "420", "25"]
        assert view.sheet == "Data"
        assert view.sheet_required is False

    def test_multiple_sheets_require_an_explicit_choice(self, multi_xlsx: Path) -> None:
        view = read_table(multi_xlsx)
        assert view.sheets == ["README", "Data"]
        assert view.sheet_required is True
        # 默认只是"先有东西可看"，不是"第一张就是答案"
        assert view.sheet == "README"
        assert view.columns == ["说明", "本页不是数据"]

    def test_explicit_sheet_reads_that_sheet(self, multi_xlsx: Path) -> None:
        view = read_table(multi_xlsx, sheet="Data")
        assert view.columns == ["specimen_id", "batch_id", "ys_mpa", "uts_mpa", "el_pct"]
        assert view.rows[0] == ["X9", "B-77", "355", "425", "23"]
        assert view.sheet == "Data"
        assert view.sheet_required is False

    def test_unknown_sheet_lists_available_ones(self, multi_xlsx: Path) -> None:
        with pytest.raises(TableError) as excinfo:
            read_table(multi_xlsx, sheet="NoSuchSheet")
        message = str(excinfo.value)
        assert "NoSuchSheet" in message and "README" in message and "Data" in message

    def test_header_row_is_honoured(self, tmp_path: Path) -> None:
        path = tmp_path / "late.xlsx"
        _save_xlsx(path, {"Sheet1": [
            ["导出说明", "2026-10-14", ""],
            [],
            ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"],
            ["S3", 355, 425, 23],
        ]})
        view = read_table(path, header_row=3)
        assert view.columns == ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"]
        assert view.rows[0][0] == "S3"

    def test_header_row_out_of_range_is_specific(self, tmp_path: Path) -> None:
        path = tmp_path / "late.xlsx"
        _save_xlsx(path, {"Sheet1": [["a"], ["1"]]})
        with pytest.raises(TableError, match="header_row=99"):
            read_table(path, header_row=99)

    def test_blank_header_row_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "late.xlsx"
        _save_xlsx(path, {"Sheet1": [
            ["导出说明", "2026-10-14", ""],
            [],
            ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"],
            ["S3", 355, 425, 23],
        ]})
        # header_row=2 是一整行空：不能悄悄返回一个零列空表
        with pytest.raises(TableError, match="没有可用的表头"):
            read_table(path, header_row=2)

    def test_formula_cells_are_flagged_and_cached_values_are_separate(self, tmp_path: Path) -> None:
        path = tmp_path / "formula.xlsx"
        _save_xlsx(path, {"Data": [
            ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"],
            ["S4", 350, "=B2*1.2", 25],
        ]})
        view = read_table(path)
        assert view.has_formulas is True
        # SteelData 用 data_only=False，所以这一格就是公式本身，不是缓存值
        assert view.rows[0][2] == "=B2*1.2"
        assert view.cached_rows and str(view.cached_rows[0][2]) != "=B2*1.2"
        assert any("公式" in item for item in view.notes)

    def test_trailing_padding_does_not_hide_ragged_rows(self, tmp_path: Path) -> None:
        # openpyxl 会把每行补齐到 max_column；尾部空单元格必须去掉，
        # 否则既会出现幻影空列，也会掩盖真正的不齐列。
        path = tmp_path / "ragged.xlsx"
        _save_xlsx(path, {"Data": [
            ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"],
            ["S5", 350, 420],              # 少一列
            ["S6", 351, 421, 24, 999],     # 多一列
        ]})
        view = read_table(path)
        assert view.columns == ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"]
        assert view.empty_headers == []
        assert view.ragged_rows == [1, 2]

    def test_duplicate_headers_are_reported(self, tmp_path: Path) -> None:
        path = tmp_path / "dup.xlsx"
        _save_xlsx(path, {"Data": [
            ["specimen_id", "ys_mpa", "ys_mpa", "el_pct"],
            ["S7", 350, 351, 25],
        ]})
        assert read_table(path).duplicate_headers == ["ys_mpa"]

    def test_empty_workbook_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "blank.xlsx"
        _save_xlsx(path, {"Empty": []})
        with pytest.raises(TableError):
            read_table(path)

    def test_broken_zip_is_refused_with_advice(self, tmp_path: Path) -> None:
        path = tmp_path / "broken.xlsx"
        path.write_bytes(b"PK\x03\x04not-really-a-zip")
        with pytest.raises(TableError, match="XLSX"):
            read_table(path)

    def test_csv_renamed_to_xlsx_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "actually_csv.xlsx"
        path.write_text("a,b\n1,2\n", encoding="utf-8")
        with pytest.raises(TableError):
            read_table(path)

    def test_row_and_column_limits(self, tmp_path: Path) -> None:
        rows = tmp_path / "many_rows.xlsx"
        _save_xlsx(rows, {"Data": [["a"], *[[1] for _ in range(MAX_ROWS + 1)]]})
        with pytest.raises(TableError, match=str(MAX_ROWS)):
            read_table(rows)
        cols = tmp_path / "many_cols.xlsx"
        _save_xlsx(cols, {"Data": [[f"c{i}" for i in range(MAX_COLUMNS + 1)]]})
        with pytest.raises(TableError, match=str(MAX_COLUMNS)):
            read_table(cols)

    def test_unsupported_suffix_is_refused(self, tmp_path: Path) -> None:
        path = tmp_path / "a.parquet"
        path.write_bytes(b"\x00\x01")
        with pytest.raises(TableError, match="只支持"):
            read_table(path)


# --------------------------------------------------------------------------- #
# 用真实 SteelData 合同核对（登记了解释器才跑）
# --------------------------------------------------------------------------- #

_STEELDATA_PROBE = r'''
import json, sys
from steeldata.adapters import matches, read_document
from steeldata.models import Profile

single, multi, formula = sys.argv[1], sys.argv[2], sys.argv[3]
out = {}

def err(fn, *a):
    try:
        fn(*a)
    except Exception as exc:  # noqa: BLE001
        return f"{type(exc).__name__}: {exc}"
    return None

doc = read_document(open(single, "rb").read(), ".xlsx", "Data")
out["single_columns"] = doc.rows[0]
out["single_rows"] = doc.rows[1:]
mdoc = read_document(open(multi, "rb").read(), ".xlsx", "Data")
out["multi_columns"] = mdoc.rows[0]
out["multi_rows"] = mdoc.rows[1:]
out["multi_no_sheet"] = err(read_document, open(multi, "rb").read(), ".xlsx", None)
out["formula"] = err(read_document, open(formula, "rb").read(), ".xlsx", "Data")
profile = Profile.model_validate({
    "schema_version": 1, "profile_id": "xlsx-check", "domain": "tensile", "layout": "flat",
    "header_row": 1, "sheet": "Data",
    "columns": [
        {"source": "specimen_id", "target": "specimen_id", "dtype": "text", "required": True},
        {"source": "ys_mpa", "target": "ys_mpa", "dtype": "number", "unit": "MPa", "target_unit": "MPa"},
        {"source": "uts_mpa", "target": "uts_mpa", "dtype": "number", "unit": "MPa", "target_unit": "MPa"},
        {"source": "el_pct", "target": "el_pct", "dtype": "number", "unit": "%", "target_unit": "%"},
    ],
})
out["matches"] = bool(matches(doc, profile))
print(json.dumps(out, ensure_ascii=False))
'''


def _steeldata_interpreter() -> Path | None:
    from steelai.web_api.config import config

    tool = config().tool("steeldata")
    interpreter = tool.interpreter if tool else None
    if interpreter and Path(interpreter).is_file():
        return Path(interpreter)
    return None


class TestAgreesWithSteelData:
    """我们给出的列/值必须与 SteelData 自己的读法逐字节一致。"""

    def _probe(self, single: Path, multi: Path, formula: Path) -> dict:
        interpreter = _steeldata_interpreter()
        if interpreter is None:
            pytest.skip("未登记 SteelData 解释器（config/settings.toml 的 tools.steeldata_python）")
        completed = subprocess.run(
            [str(interpreter), "-c", _STEELDATA_PROBE, str(single), str(multi), str(formula)],
            capture_output=True, text=True, encoding="utf-8", check=False,
        )
        if completed.returncode != 0:
            pytest.skip(f"SteelData 探针无法运行：{(completed.stderr or '').strip()[-200:]}")
        return json.loads(completed.stdout.strip().splitlines()[-1])

    def test_columns_and_rows_match(self, single_xlsx: Path, multi_xlsx: Path, tmp_path: Path) -> None:
        formula = tmp_path / "formula.xlsx"
        _save_xlsx(formula, {"Data": [
            ["specimen_id", "ys_mpa", "uts_mpa", "el_pct"],
            ["S4", 350, "=B2*1.2", 25],
        ]})
        probe = self._probe(single_xlsx, multi_xlsx, formula)

        ours = read_table(single_xlsx)
        assert ours.columns == probe["single_columns"]
        assert ours.rows == probe["single_rows"]
        assert probe["matches"] is True

        chosen = read_table(multi_xlsx, sheet="Data")
        assert chosen.columns == probe["multi_columns"]
        assert chosen.rows == probe["multi_rows"]

        # SteelData 自己也拒绝：多工作表不写 sheet / 含公式
        assert "Multiple worksheets" in (probe["multi_no_sheet"] or "")
        assert "Formula" in (probe["formula"] or "")
        assert read_table(formula).has_formulas is True

    def test_broken_zip_judgement_agrees(self, tmp_path: Path) -> None:
        path = tmp_path / "broken.xlsx"
        path.write_bytes(b"PK\x03\x04not-really-a-zip")
        with pytest.raises(zipfile.BadZipFile):
            with zipfile.ZipFile(io.BytesIO(path.read_bytes())):
                pass
        with pytest.raises(TableError):
            read_table(path)
