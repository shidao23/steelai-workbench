"""WO-31 §2/§3 — 科研项目/任务登记读取。

项目清单**直接读现有的科研登记**（科研主表 ``项目表`` 工作表），不新建第二个
总项目库。任务沿用 ``03_工作区/<project_id>/<task_id>/`` 这个已有 owner 结构，
每项包含：

```
<project>/<task_id>/
├── TASK.md          # 任务说明（人/换 AI 都能读）
├── inputs/          # 本任务上传的输入（原件只读，不改写）
├── runs/<run_id>/   # 每次运行的唯一输出目录
└── handoff.md       # 接手摘要
```

真实实验正文与 ``lab_data`` 不在这里扫描：登记表只提供项目身份，
真实数据根由被忽略的本地配置指向。
"""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

from steelai.web_api.config import config
from steelai.web_api.security import GuardError, safe_identifier

TASK_MARKER = "TASK.md"
SYNTHETIC_PROJECT_ID = "PSYN"
SYNTHETIC_PROJECT_NAME = "合成演示（本批隔离资料）"
ALLOWED_INPUT_SUFFIXES = {
    ".csv", ".tsv", ".txt", ".json", ".md",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp",
    ".xlsx", ".xls", ".yaml", ".yml",
}
MAX_UPLOAD_BYTES = 64 * 1024 * 1024


@dataclass(frozen=True)
class Project:
    project_id: str
    directory_name: str
    short_name: str
    status: str
    source: str
    synthetic: bool


@dataclass(frozen=True)
class Task:
    project_id: str
    task_id: str
    title: str
    created_at: str
    run_count: int
    file_count: int


def _master_projects() -> tuple[list[Project], str]:
    """Read the ``项目表`` sheet of the existing research master workbook."""
    cfg = config()
    path = cfg.master_workbook
    if not path.is_file():
        return [], f"科研主表不存在：{path}"
    try:
        import openpyxl  # noqa: PLC0415 - optional at import time
    except ImportError:
        return [], "读取科研主表需要 openpyxl（未安装）"
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - report, never crash the API
        return [], f"科研主表无法打开：{exc}"
    if "项目表" not in workbook.sheetnames:
        workbook.close()
        return [], "科研主表缺少“项目表”工作表"
    sheet = workbook["项目表"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        workbook.close()
        return [], "“项目表”为空"
    index = {str(name).strip(): position for position, name in enumerate(header) if name}
    projects: list[Project] = []
    for row in rows:
        if not row:
            continue
        pid = str(row[index.get("项目ID", 0)] or "").strip()
        if not re.fullmatch(r"P\d{2,4}", pid):
            continue
        def cell(key: str, fallback: int) -> str:
            position = index.get(key, fallback)
            return str(row[position] or "").strip() if position < len(row) else ""
        projects.append(Project(
            project_id=pid,
            directory_name=cell("项目目录名", 1),
            short_name=cell("项目简称", 2),
            status=cell("项目状态", 5),
            source=cell("来源路径", 7),
            synthetic=False,
        ))
    workbook.close()
    return projects, ""


def list_projects() -> tuple[list[Project], list[str]]:
    projects, problem = _master_projects()
    notes: list[str] = []
    if problem:
        notes.append(problem)
    cfg = config()
    if cfg.synthetic_root.is_dir():
        projects.append(Project(
            project_id=SYNTHETIC_PROJECT_ID,
            directory_name=cfg.synthetic_root.name,
            short_name=SYNTHETIC_PROJECT_NAME,
            status="合成",
            source="本批合成资料（隔离，不读真实实验正文）",
            synthetic=True,
        ))
    return projects, notes


def find_project(project_id: str) -> Project:
    pid = safe_identifier(project_id, "项目标识")
    projects, _ = list_projects()
    for project in projects:
        if project.project_id == pid:
            return project
    raise GuardError(f"未登记的项目：{pid}")


def project_root(project_id: str) -> Path:
    project = find_project(project_id)
    cfg = config()
    if project.synthetic:
        return cfg.synthetic_root
    root = cfg.work_root / project.project_id
    return root


def ensure_project_root(project_id: str) -> Path:
    root = project_root(project_id)
    root.mkdir(parents=True, exist_ok=True)
    return root


def task_root(project_id: str, task_id: str) -> Path:
    tid = safe_identifier(task_id, "任务标识")
    root = project_root(project_id) / tid
    return root


def list_tasks(project_id: str) -> list[Task]:
    root = project_root(project_id)
    if not root.is_dir():
        return []
    tasks: list[Task] = []
    for child in sorted(root.iterdir()):
        if not child.is_dir() or not (child / TASK_MARKER).is_file():
            continue
        runs = child / "runs"
        run_count = sum(1 for item in runs.iterdir() if item.is_dir()) if runs.is_dir() else 0
        inputs = child / "inputs"
        file_count = sum(1 for item in inputs.iterdir() if item.is_file()) if inputs.is_dir() else 0
        title = ""
        try:
            title = (child / TASK_MARKER).read_text(encoding="utf-8").splitlines()[0].lstrip("# ").strip()
        except (OSError, IndexError):
            title = child.name
        created = datetime.fromtimestamp(child.stat().st_mtime).isoformat(timespec="seconds")
        tasks.append(Task(project_id, child.name, title or child.name, created, run_count, file_count))
    return tasks


def create_task(project_id: str, task_id: str, title: str, note: str = "") -> Task:
    tid = safe_identifier(task_id, "任务标识")
    root = ensure_project_root(project_id)
    target = root / tid
    if target.exists():
        raise GuardError(f"任务已存在：{tid}")
    (target / "inputs").mkdir(parents=True)
    (target / "runs").mkdir(parents=True)
    heading = (title or tid).strip().replace("\n", " ")
    body = [f"# {heading}", "", f"- project_id: `{project_id}`", f"- task_id: `{tid}`",
            f"- created: {datetime.now().isoformat(timespec='seconds')}", ""]
    if note:
        body += ["## 说明", "", note.strip(), ""]
    body += ["## 约定", "",
             "- 输入只读，运行输出写 `runs/<run_id>/`，不覆盖既有运行。",
             "- 未知单位/身份/摘要变化必须停在 `needs_human`，不自动改绿。",
             "- 候选成果不自动晋升；正式晋升走原 owner 人审。", ""]
    (target / TASK_MARKER).write_text("\n".join(body), encoding="utf-8")
    (target / "handoff.md").write_text(
        f"# 接手摘要 · {heading}\n\n- 尚未运行。\n", encoding="utf-8")
    return Task(project_id, tid, heading, datetime.now().isoformat(timespec="seconds"), 0, 0)


def task_context(project_id: str, task_id: str) -> dict:
    root = task_root(project_id, task_id)
    if not (root / TASK_MARKER).is_file():
        raise GuardError(f"任务不存在：{project_id}/{task_id}")
    task_md = (root / TASK_MARKER).read_text(encoding="utf-8")
    handoff = ""
    handoff_path = root / "handoff.md"
    if handoff_path.is_file():
        handoff = handoff_path.read_text(encoding="utf-8")
    return {
        "project_id": project_id,
        "task_id": task_id,
        "root": str(root),
        "task_markdown": task_md,
        "handoff_markdown": handoff,
        "inputs": [item.name for item in sorted((root / "inputs").iterdir()) if item.is_file()],
    }


def list_input_files(project_id: str, task_id: str) -> list[dict]:
    root = task_root(project_id, task_id) / "inputs"
    if not root.is_dir():
        return []
    rows: list[dict] = []
    for item in sorted(root.iterdir()):
        if not item.is_file():
            continue
        stat = item.stat()
        rows.append({
            "file_id": item.name,
            "name": item.name,
            "size": stat.st_size,
            "suffix": item.suffix.lower(),
            "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            "previewable": item.suffix.lower() in ALLOWED_INPUT_SUFFIXES,
        })
    return rows


def input_path(project_id: str, task_id: str, file_id: str) -> Path:
    from steelai.web_api.security import safe_filename, resolve_under

    root = task_root(project_id, task_id) / "inputs"
    name = safe_filename(file_id)
    return resolve_under(root, name)


def register_upload(project_id: str, task_id: str, filename: str, payload: bytes) -> dict:
    from steelai.web_api.security import safe_filename

    if len(payload) > MAX_UPLOAD_BYTES:
        raise GuardError(f"上传超过上限 {MAX_UPLOAD_BYTES} 字节")
    name = safe_filename(filename)
    if Path(name).suffix.lower() not in ALLOWED_INPUT_SUFFIXES:
        raise GuardError(f"不接受该类型：{Path(name).suffix or '（无扩展名）'}")
    root = task_root(project_id, task_id) / "inputs"
    if not root.is_dir():
        raise GuardError("任务不存在或缺少 inputs 目录")
    target = root / name
    if target.exists():
        raise GuardError(f"同名输入已存在，先改名或删除：{name}")
    target.write_bytes(payload)
    return {"file_id": name, "name": name, "size": len(payload)}


def delete_input(project_id: str, task_id: str, file_id: str) -> None:
    target = input_path(project_id, task_id, file_id)
    if not target.is_file():
        raise GuardError("文件不存在")
    target.unlink()


def project_summary(project: Project) -> dict:
    root = project_root(project.project_id)
    tasks = list_tasks(project.project_id)
    return {
        **asdict(project),
        "root_exists": root.is_dir(),
        "task_count": len(tasks),
        "tasks": [asdict(task) for task in tasks],
    }


def reset_synthetic(project_id: str) -> dict:
    """Clear the isolated synthetic project only (never a real project)."""
    if project_id != SYNTHETIC_PROJECT_ID:
        raise GuardError("只允许重置合成演示项目")
    root = config().synthetic_root
    if root.is_dir():
        shutil.rmtree(root)
    root.mkdir(parents=True, exist_ok=True)
    return {"reset": True, "root": str(root)}


def write_handoff(project_id: str, task_id: str, markdown: str) -> Path:
    root = task_root(project_id, task_id)
    if not (root / TASK_MARKER).is_file():
        raise GuardError("任务不存在")
    path = root / "handoff.md"
    path.write_text(markdown, encoding="utf-8")
    return path


def read_json(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
