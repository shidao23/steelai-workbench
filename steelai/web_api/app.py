"""WO-31 §3 — 科研薄 API（FastAPI）。

一个进程同时提供：

* ``/api/*``：项目/任务/文件/运行/产物/能力的 JSON 接口；
* ``/``：已构建的 React 单页（``frontend/dist``，同源提供，避免跨源与凭据转发）；
* ``/research``：科研专用 Streamlit 反代（可选，见 §4）。

写入型请求校验可信 ``Origin``；所有路径/文件名/标识都经过 :mod:`security` 守卫。
"""

from __future__ import annotations

import json
import struct
from pathlib import Path

from fastapi import Body, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles

from steelai.paths import ROOT
from steelai.web_api import registry, runner, tools
from steelai.web_api.config import config
from steelai.web_api.security import GuardError, check_origin, interpreter_is_trusted

FRONTEND_DIST = ROOT / "frontend" / "dist"
PREVIEW_ROWS = 25

app = FastAPI(title="SteelAI 科研薄 API", version="0.1.0", docs_url="/api/docs",
              openapi_url="/api/openapi.json")

app.add_middleware(
    CORSMiddleware,
    allow_origins=list(config().allowed_origins),
    allow_credentials=False,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Content-Type"],
)


def _guard(callable_):
    try:
        return callable_()
    except GuardError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.middleware("http")
async def origin_guard(request: Request, call_next):
    if request.method in {"POST", "DELETE", "PUT", "PATCH"}:
        if not check_origin(request.headers.get("origin"), config().allowed_origins):
            return JSONResponse(status_code=403, content={"detail": "不可信来源"})
    return await call_next(request)


# --------------------------------------------------------------------------- #
# 能力与健康
# --------------------------------------------------------------------------- #

@app.get("/api/health")
def health() -> dict:
    cfg = config()
    return {
        "status": "ok",
        "workspace_root": str(cfg.workspace_root),
        "work_root": str(cfg.work_root),
        "run_root": str(cfg.run_root),
        "allowed_origins": list(cfg.allowed_origins),
        "research_app_url": cfg.streamlit_url,
        "frontend_built": FRONTEND_DIST.is_dir(),
        # 解释器只来自服务端解析（环境变量/被忽略的 config/settings.toml），
        # 客户端永远不能提交。这里如实报告它是否落在已登记的可信根内。
        "tools": {
            tool_id: {
                "interpreter": str(item.interpreter) if item.interpreter else None,
                "registered": item.interpreter is not None,
                "trusted": interpreter_is_trusted(item.interpreter, cfg.trusted_tool_roots),
            }
            for tool_id, item in cfg.tools.items()
        },
    }


@app.get("/api/capabilities")
def capabilities() -> dict:
    return {"tools": tools.capabilities()}


# --------------------------------------------------------------------------- #
# 项目 / 任务
# --------------------------------------------------------------------------- #

@app.get("/api/projects")
def get_projects() -> dict:
    projects, notes = registry.list_projects()
    return {"projects": [registry.project_summary(item) for item in projects], "notes": notes}


@app.get("/api/projects/{project_id}/tasks")
def get_tasks(project_id: str) -> dict:
    tasks = _guard(lambda: registry.list_tasks(project_id))
    return {"project_id": project_id, "tasks": [item.__dict__ for item in tasks]}


@app.post("/api/projects/{project_id}/tasks")
def post_task(project_id: str, payload: dict = Body(...)) -> dict:
    task = _guard(lambda: registry.create_task(
        project_id,
        str(payload.get("task_id", "")).strip(),
        str(payload.get("title", "")).strip(),
        str(payload.get("note", "")),
    ))
    return {"task": task.__dict__}


@app.get("/api/projects/{project_id}/tasks/{task_id}/context")
def get_context(project_id: str, task_id: str) -> dict:
    return _guard(lambda: registry.task_context(project_id, task_id))


@app.get("/api/projects/{project_id}/tasks/{task_id}/files")
def get_files(project_id: str, task_id: str) -> dict:
    files = _guard(lambda: registry.list_input_files(project_id, task_id))
    return {"files": files}


@app.post("/api/projects/{project_id}/tasks/{task_id}/upload")
async def post_upload(project_id: str, task_id: str, file: UploadFile = File(...)) -> dict:
    payload = await file.read()
    result = _guard(lambda: registry.register_upload(project_id, task_id, file.filename or "", payload))
    return {"uploaded": result}


@app.delete("/api/projects/{project_id}/tasks/{task_id}/files/{file_id}")
def delete_file(project_id: str, task_id: str, file_id: str) -> dict:
    _guard(lambda: registry.delete_input(project_id, task_id, file_id))
    return {"deleted": file_id}


@app.get("/api/projects/{project_id}/tasks/{task_id}/preview")
def get_preview(project_id: str, task_id: str, file_id: str) -> dict:
    path = _guard(lambda: registry.input_path(project_id, task_id, file_id))
    return _guard(lambda: _preview(path))


@app.get("/api/projects/{project_id}/tasks/{task_id}/handoff")
def get_handoff(project_id: str, task_id: str) -> dict:
    context = _guard(lambda: registry.task_context(project_id, task_id))
    return {"handoff_markdown": context["handoff_markdown"]}


@app.post("/api/projects/{project_id}/tasks/{task_id}/handoff")
def post_handoff(project_id: str, task_id: str, payload: dict = Body(...)) -> dict:
    markdown = str(payload.get("markdown", ""))
    if not markdown.strip():
        raise HTTPException(status_code=400, detail="接手摘要不能为空")
    _guard(lambda: registry.write_handoff(project_id, task_id, markdown))
    return {"written": True}


# --------------------------------------------------------------------------- #
# 运行
# --------------------------------------------------------------------------- #

@app.post("/api/projects/{project_id}/tasks/{task_id}/runs")
def post_run(project_id: str, task_id: str, payload: dict = Body(...)) -> dict:
    tool_id = str(payload.get("tool_id", "")).strip()
    params = payload.get("params") or {}
    if not isinstance(params, dict):
        raise HTTPException(status_code=400, detail="params 必须是对象")
    record = _guard(lambda: runner.create_run(project_id, task_id, tool_id, params))
    return {"run": record}


@app.get("/api/projects/{project_id}/tasks/{task_id}/runs")
def get_task_runs(project_id: str, task_id: str) -> dict:
    return {"runs": runner.list_runs(project_id, task_id)}


@app.get("/api/runs/{run_id}")
def get_run(run_id: str) -> dict:
    return {"run": _guard(lambda: runner.load_record(run_id).as_dict())}


@app.get("/api/runs/{run_id}/artifacts/{rel:path}")
def get_artifact(run_id: str, rel: str):
    path, meta = _guard(lambda: runner.artifact_for_download(run_id, rel))
    return FileResponse(path, filename=path.name, media_type=_media_type(path))


@app.get("/api/projects/{project_id}/tasks/{task_id}/artifacts")
def get_task_artifacts(project_id: str, task_id: str) -> dict:
    rows: list[dict] = []
    for record in runner.list_runs(project_id, task_id):
        for artifact in record.get("artifacts", []):
            rows.append({
                **artifact,
                "run_id": record["run_id"],
                "tool_id": record["tool_id"],
                "tool_label": record.get("tool_label", record["tool_id"]),
                "status": record["status"],
                "created": record["created"],
            })
    return {"artifacts": rows}


@app.post("/api/projects/{project_id}/tasks/{task_id}/promote")
def post_promote(project_id: str, task_id: str, payload: dict = Body(...)) -> dict:
    """候选成果晋升：必须具名人工审批，且要接上原 owner 的晋升流程。"""
    reviewer = str(payload.get("reviewer", "")).strip()
    run_id = str(payload.get("run_id", "")).strip()
    if not reviewer or reviewer.lower() in {"system", "agent", "auto", "automatic", "web"}:
        raise HTTPException(status_code=400, detail="晋升需要具名人工审批人，不能用系统/代理名")
    record = _guard(lambda: runner.load_record(run_id))
    if record.status != "succeeded":
        raise HTTPException(status_code=409, detail="只有成功的运行才能晋升")
    return JSONResponse(status_code=501, content={
        "detail": "晋升尚未接通：原 owner 的正式成果晋升流程未在本单接线。",
        "status": "not_connected",
        "candidate": {
            "run_id": run_id,
            "artifacts": record.artifacts,
            "requested_reviewer": reviewer,
        },
        "note": "候选仍保留在任务 runs/ 下，不会被自动 accepted 或覆盖 CURRENT。",
    })


# --------------------------------------------------------------------------- #
# 预览实现（只读，不改写输入）
# --------------------------------------------------------------------------- #

def _media_type(path: Path) -> str:
    return {
        ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
        ".tif": "image/tiff", ".tiff": "image/tiff", ".bmp": "image/bmp",
        ".html": "text/html", ".json": "application/json", ".csv": "text/csv",
        ".txt": "text/plain", ".md": "text/markdown", ".npy": "application/octet-stream",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }.get(path.suffix.lower(), "application/octet-stream")


def _image_size(path: Path) -> dict:
    """Read pixel dimensions from the file header without decoding the image."""
    with path.open("rb") as handle:
        head = handle.read(32)
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            width, height = struct.unpack(">II", head[16:24])
            return {"format": "png", "width": width, "height": height}
        if head[:2] == b"\xff\xd8":
            handle.seek(2)
            while True:
                marker = handle.read(2)
                if len(marker) < 2 or marker[0] != 0xFF:
                    break
                length = struct.unpack(">H", handle.read(2))[0]
                if marker[1] in (0xC0, 0xC1, 0xC2, 0xC3):
                    data = handle.read(5)
                    height, width = struct.unpack(">HH", data[1:5])
                    return {"format": "jpeg", "width": width, "height": height}
                handle.seek(length - 2, 1)
        if head[:2] == b"BM":
            width, height = struct.unpack("<ii", head[18:26])
            return {"format": "bmp", "width": abs(width), "height": abs(height)}
    return {"format": path.suffix.lstrip(".").lower(), "width": None, "height": None}


def _preview(path: Path) -> dict:
    if not path.is_file():
        raise GuardError("文件不存在")
    suffix = path.suffix.lower()
    base = {"file_id": path.name, "size": path.stat().st_size, "suffix": suffix}
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        return {**base, "kind": "image", "image": _image_size(path),
                "note": "仅读取文件头尺寸，不解码、不修改输入。"}
    if suffix in {".csv", ".tsv"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        lines = [line for line in text.splitlines() if line.strip()]
        delimiter = "\t" if suffix == ".tsv" else _guess_delimiter(lines[0] if lines else "")
        rows = [line.split(delimiter) for line in lines[:PREVIEW_ROWS + 1]]
        return {**base, "kind": "table", "delimiter": delimiter,
                "columns": rows[0] if rows else [], "rows": rows[1:],
                "total_lines": len(lines), "truncated": len(lines) > PREVIEW_ROWS + 1}
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig", errors="replace"))
        return {**base, "kind": "json", "json_head": _truncate_json(data)}
    if suffix in {".xlsx", ".xls"}:
        try:
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            rows = [["" if cell is None else str(cell) for cell in row]
                    for row in sheet.iter_rows(max_row=PREVIEW_ROWS + 1, values_only=True)]
            sheets = list(workbook.sheetnames)
            workbook.close()
            return {**base, "kind": "workbook", "sheets": sheets, "sheet": sheets[0],
                    "columns": rows[0] if rows else [], "rows": rows[1:]}
        except Exception as exc:  # noqa: BLE001
            return {**base, "kind": "binary", "note": f"工作簿无法预览：{exc}"}
    if suffix in {".txt", ".md", ".yaml", ".yml"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        return {**base, "kind": "text", "text": text[:4000],
                "truncated": len(text) > 4000}
    return {**base, "kind": "binary", "note": "该类型不提供文本预览，可直接下载。"}


def _guess_delimiter(header: str) -> str:
    counts = {",": header.count(","), ";": header.count(";"), "\t": header.count("\t")}
    return max(counts, key=counts.get) if max(counts.values()) else ","


def _truncate_json(value, depth: int = 0):
    if depth > 3:
        return "…"
    if isinstance(value, dict):
        return {key: _truncate_json(item, depth + 1) for key, item in list(value.items())[:20]}
    if isinstance(value, list):
        return [_truncate_json(item, depth + 1) for item in value[:10]]
    return value


# --------------------------------------------------------------------------- #
# 前端静态资源（同源提供，避免跨源与凭据转发）
# --------------------------------------------------------------------------- #

if FRONTEND_DIST.is_dir():
    app.mount("/", StaticFiles(directory=str(FRONTEND_DIST), html=True), name="frontend")
